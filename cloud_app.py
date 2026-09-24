from __future__ import annotations

import csv
import hmac
import io
import json
import os
import re
import threading
import time
import uuid
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo
from typing import Any

from flask import Flask, Response, jsonify, render_template, request
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import requests

app = Flask(__name__)
app.config["JSON_AS_ASCII"] = False

_lock = threading.Lock()
_STATE_FILE = Path(os.getenv("CLOUD_STATE_FILE", "/tmp/denni_pov_state.json"))
_DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
_KOSTKA_URL = "https://api.dataovozidlech.cz/api/vehicletechnicaldata/v2"
_KOSTKA_VIN = re.compile(r"[A-HJ-NPR-Z0-9]{17}\Z")
_kostka_worker_lock = threading.Lock()
_KOSTKA_REFRESH_DAYS = 30
_KOSTKA_MAX_PER_CHECK = 15


class KostkaLimitReached(RuntimeError):
    """The provider asked us to pause requests for this API key."""


def _kostka_keys() -> list[str]:
    return list(dict.fromkeys(key for key in (
        os.getenv("DATOVA_KOSTKA_API_KEY", "").strip(),
        os.getenv("DATOVA_KOSTKA_API_KEY_1", "").strip(),
        os.getenv("DATOVA_KOSTKA_API_KEY_2", "").strip(),
    ) if key))


def _kostka_saved(vin: str) -> tuple[dict[str, Any], str] | None:
    if not _DATABASE_URL:
        return None
    with _db_connect() as conn:
        with conn.cursor() as cur:
            _db_init(cur)
            cur.execute("SELECT data, fetched_at FROM denni_pov_vehicle_technical WHERE vin = %s", (vin,))
            row = cur.fetchone()
            conn.commit()
    return (row[0], row[1].isoformat()) if row else None


def _kostka_fetch(vin: str, keys: list[str]) -> dict[str, Any] | None:
    for key in keys:
        response = requests.get(_KOSTKA_URL, params={"vin": vin},
                                headers={"api_key": key}, timeout=12)
        if response.status_code in (401, 403):
            continue
        if response.status_code in (429, 503) or response.status_code >= 500:
            raise KostkaLimitReached("Datová kostka požaduje pauzu mezi dotazy.")
        response.raise_for_status()
        payload = response.json()
        data = payload.get("Data") if isinstance(payload, dict) else None
        return data if isinstance(payload, dict) and payload.get("Status") == 1 and isinstance(data, dict) and data else None
    raise ValueError("API klíč Datové kostky nebyl přijat.")


def _kostka_update(vins: list[str]) -> None:
    if not _kostka_worker_lock.acquire(blocking=False):
        return
    try:
        keys = _kostka_keys()
        if not keys or not _DATABASE_URL:
            return
        with _db_connect() as conn:
            with conn.cursor() as cur:
                _db_init(cur)
                cur.execute("""
                    SELECT vin FROM denni_pov_vehicle_technical
                    WHERE vin = ANY(%s) AND fetched_at >= NOW() - (%s * INTERVAL '1 day')
                """, (vins, _KOSTKA_REFRESH_DAYS))
                fresh = {row[0] for row in cur.fetchall()}
                conn.commit()
        attempted = 0
        for vin in dict.fromkeys(vins):
            if vin in fresh:
                continue
            if attempted >= _KOSTKA_MAX_PER_CHECK:
                break
            attempted += 1
            try:
                technical = _kostka_fetch(vin, keys)
                with _db_connect() as conn:
                    with conn.cursor() as cur:
                        cur.execute("""
                            INSERT INTO denni_pov_vehicle_technical (vin, data, fetched_at)
                            VALUES (%s, %s::jsonb, NOW())
                            ON CONFLICT (vin) DO UPDATE
                            SET data = EXCLUDED.data, fetched_at = EXCLUDED.fetched_at
                        """, (vin, json.dumps(technical or {}, ensure_ascii=False)))
                        conn.commit()
            except KostkaLimitReached:
                print("Datová kostka: limit požadavků, další VIN počkají na následující kontrolu.")
                break
            except (requests.RequestException, ValueError, PersistenceUnavailable) as exc:
                print(f"Datová kostka {vin}: {exc.__class__.__name__}")
                if isinstance(exc, (ValueError, PersistenceUnavailable)):
                    break
            time.sleep(1)
    except Exception as exc:
        print(f"Automatické načtení Datové kostky selhalo: {exc.__class__.__name__}")
    finally:
        _kostka_worker_lock.release()


def _kostka_start(rows: list[dict[str, Any]]) -> None:
    vins = [vin for row in rows if isinstance(row, dict)
            if _KOSTKA_VIN.fullmatch(vin := re.sub(r"\s+", "", str(row.get("vin") or "")).upper())]
    if vins and _DATABASE_URL and _kostka_keys():
        threading.Thread(target=_kostka_update, args=(vins,), daemon=True,
                         name="datova-kostka-refresh").start()

class PersistenceUnavailable(RuntimeError):
    """Raised when durable annotation storage is not available."""


def _safe_db_error(exc: Exception) -> str:
    message = str(exc).strip() or exc.__class__.__name__
    if _DATABASE_URL:
        message = message.replace(_DATABASE_URL, "[DATABASE_URL]")
    return message[:500]


def _db_connect():
    if not _DATABASE_URL:
        raise PersistenceUnavailable("DATABASE_URL není na Renderu nastavená.")
    try:
        import psycopg
        return psycopg.connect(_DATABASE_URL, connect_timeout=8)
    except Exception as exc:
        raise PersistenceUnavailable(_safe_db_error(exc)) from exc


def _db_probe() -> tuple[bool, str]:
    try:
        with _db_connect() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT 1")
                cur.fetchone()
        return True, ""
    except PersistenceUnavailable as exc:
        return False, str(exc)


def _normalize_annotations(source: dict[str, Any] | None) -> dict[str, dict[str, str]]:
    source = source if isinstance(source, dict) else {}
    return {
        str(key).strip().upper(): {
            "note": str((value or {}).get("note") or ""),
            "workflow_status": str((value or {}).get("workflow_status") or ""),
            "updated_at": str((value or {}).get("updated_at") or ""),
        }
        for key, value in source.items()
        if isinstance(value, dict) and str(key or "").strip()
    }


_PRAGUE_TZ = ZoneInfo("Europe/Prague")


def _now_dt() -> datetime:
    return datetime.now(_PRAGUE_TZ)


def _now() -> str:
    return _now_dt().strftime("%d.%m.%Y %H:%M:%S")


def _today_iso() -> str:
    return _now_dt().date().isoformat()


def _date_from_display(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        return datetime.strptime(text[:10], "%d.%m.%Y").date().isoformat()
    except ValueError:
        return ""


def _default_state() -> dict[str, Any]:
    return {
        "running": False,
        "started_at": None,
        "finished_at": None,
        "error": None,
        "progress": {"percent": 0, "phase": "Připraveno", "eta_seconds": 0},
        "sources": {
            "tirbazar": {"state": "idle", "status": "Čekám na synchronizaci", "detail": "Kancelářský agent / TIRBazar SQL"},
            "uniqa": {"state": "idle", "status": "Čekám na synchronizaci", "detail": "AIV / Denní POV / Aktivní"},
            "allianz": {"state": "idle", "status": "Čekám na synchronizaci", "detail": "Flotilové PDF"},
        },
        "summary": {"active": 0, "ok_total": 0, "ok_uniqa": 0, "ok_allianz": 0, "missing": 0, "absent_insured": 0, "absent_uninsured": 0, "deposit": 0, "sold_uniqa": 0, "extra_uniqa": 0},
        "results": [],
        "changes": {"count": 0, "items": [], "summary_delta": {}, "compared_to": None},
        "annotations": {},
        "csv_available": False,
        "synced_at": None,
        "_command": None,
    }


def _db_init(cur) -> None:
    cur.execute("""
        CREATE TABLE IF NOT EXISTS denni_pov_state (
            id INTEGER PRIMARY KEY,
            payload JSONB NOT NULL,
            updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS denni_pov_history (
            id BIGSERIAL PRIMARY KEY,
            checked_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            started_at TEXT,
            finished_at TEXT,
            summary JSONB NOT NULL,
            results JSONB NOT NULL,
            error TEXT,
            alert_sent BOOLEAN NOT NULL DEFAULT FALSE,
            changes JSONB NOT NULL DEFAULT '{}'::jsonb,
            daily_date DATE
        )
    """)
    cur.execute("ALTER TABLE denni_pov_history ADD COLUMN IF NOT EXISTS changes JSONB NOT NULL DEFAULT '{}'::jsonb")
    cur.execute("ALTER TABLE denni_pov_history ADD COLUMN IF NOT EXISTS daily_date DATE")
    cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS denni_pov_history_daily_date_uidx ON denni_pov_history (daily_date) WHERE daily_date IS NOT NULL")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS denni_pov_annotations (
            vehicle_key TEXT PRIMARY KEY,
            note TEXT NOT NULL DEFAULT '',
            workflow_status TEXT NOT NULL DEFAULT '',
            updated_at TEXT NOT NULL
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS denni_pov_annotation_history (
            id BIGSERIAL PRIMARY KEY,
            vehicle_key TEXT NOT NULL,
            vin TEXT NOT NULL DEFAULT '',
            spz TEXT NOT NULL DEFAULT '',
            vehicle TEXT NOT NULL DEFAULT '',
            original_status TEXT NOT NULL DEFAULT '',
            old_workflow_status TEXT NOT NULL DEFAULT '',
            new_workflow_status TEXT NOT NULL DEFAULT '',
            old_note TEXT NOT NULL DEFAULT '',
            new_note TEXT NOT NULL DEFAULT '',
            changed_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            changed_at_text TEXT NOT NULL DEFAULT ''
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS denni_pov_annotation_history_key_idx ON denni_pov_annotation_history (vehicle_key, id DESC)")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS denni_pov_vehicle_technical (
            vin TEXT PRIMARY KEY,
            data JSONB NOT NULL,
            fetched_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
    """)


def _db_load() -> dict[str, Any] | None:
    if not _DATABASE_URL:
        return None
    try:
        with _db_connect() as conn:
            with conn.cursor() as cur:
                _db_init(cur)
                cur.execute("SELECT payload FROM denni_pov_state WHERE id = 1")
                row = cur.fetchone()
                conn.commit()
                if row:
                    payload = row[0]
                    if isinstance(payload, str):
                        payload = json.loads(payload)
                    if isinstance(payload, dict):
                        return payload
    except Exception as exc:
        print(f"Render Postgres load fallback: {exc}")
    return None


def _db_save(data: dict[str, Any]) -> bool:
    if not _DATABASE_URL:
        return False
    try:
        with _db_connect() as conn:
            with conn.cursor() as cur:
                _db_init(cur)
                cur.execute("""
                    INSERT INTO denni_pov_state (id, payload, updated_at)
                    VALUES (1, %s::jsonb, NOW())
                    ON CONFLICT (id)
                    DO UPDATE SET payload = EXCLUDED.payload, updated_at = NOW()
                """, (json.dumps(data, ensure_ascii=False),))
                conn.commit()
        return True
    except Exception as exc:
        print(f"Render Postgres save fallback: {exc}")
        return False


def _merge_change_items(existing: list[dict[str, Any]], incoming: list[dict[str, Any]]) -> list[dict[str, Any]]:
    merged = list(existing or [])
    seen = {
        (
            str(item.get("type") or ""),
            str(item.get("key") or ""),
            str(item.get("old_status") or item.get("old_spz") or ""),
            str(item.get("new_status") or item.get("new_spz") or ""),
        )
        for item in merged
    }
    for item in incoming or []:
        signature = (
            str(item.get("type") or ""),
            str(item.get("key") or ""),
            str(item.get("old_status") or item.get("old_spz") or ""),
            str(item.get("new_status") or item.get("new_spz") or ""),
        )
        if signature not in seen:
            clean = dict(item)
            clean["detected_at"] = clean.get("detected_at") or _now()
            merged.append(clean)
            seen.add(signature)
    return merged


def _daily_history_existing(day_iso: str) -> dict[str, Any] | None:
    if not _DATABASE_URL:
        return None
    try:
        with _db_connect() as conn:
            with conn.cursor() as cur:
                _db_init(cur)
                cur.execute("""
                    SELECT id, changes, started_at, finished_at
                    FROM denni_pov_history
                    WHERE daily_date = %s
                    LIMIT 1
                """, (day_iso,))
                row = cur.fetchone()
                conn.commit()
        if not row:
            return None
        return {
            "id": row[0],
            "changes": row[1] if isinstance(row[1], dict) else {},
            "started_at": row[2] or "",
            "finished_at": row[3] or "",
        }
    except Exception as exc:
        print(f"Daily history load failed: {exc}")
        return None


def _history_save_daily(data: dict[str, Any], alert_sent: bool) -> dict[str, Any]:
    day_iso = _today_iso()
    existing = _daily_history_existing(day_iso) or {}
    old_changes = existing.get("changes") if isinstance(existing.get("changes"), dict) else {}
    new_changes = data.get("changes") if isinstance(data.get("changes"), dict) else {}
    merged_items = _merge_change_items(old_changes.get("items") or [], new_changes.get("items") or [])
    daily_changes = {
        "count": len(merged_items),
        "items": merged_items,
        "summary_delta": new_changes.get("summary_delta") or old_changes.get("summary_delta") or {},
        "compared_to": old_changes.get("compared_to") or new_changes.get("compared_to"),
        "day": day_iso,
    }
    data["changes"] = daily_changes

    if not _DATABASE_URL:
        return daily_changes
    try:
        with _db_connect() as conn:
            with conn.cursor() as cur:
                _db_init(cur)
                cur.execute("""
                    INSERT INTO denni_pov_history
                        (daily_date, started_at, finished_at, summary, results, error, alert_sent, changes)
                    VALUES (%s, %s, %s, %s::jsonb, %s::jsonb, %s, %s, %s::jsonb)
                    ON CONFLICT (daily_date) WHERE daily_date IS NOT NULL
                    DO UPDATE SET
                        checked_at = NOW(),
                        started_at = COALESCE(denni_pov_history.started_at, EXCLUDED.started_at),
                        finished_at = EXCLUDED.finished_at,
                        summary = EXCLUDED.summary,
                        results = EXCLUDED.results,
                        error = EXCLUDED.error,
                        alert_sent = denni_pov_history.alert_sent OR EXCLUDED.alert_sent,
                        changes = EXCLUDED.changes
                """, (
                    day_iso,
                    existing.get("started_at") or data.get("started_at"),
                    data.get("finished_at"),
                    json.dumps(data.get("summary") or {}, ensure_ascii=False),
                    json.dumps(data.get("results") or [], ensure_ascii=False),
                    data.get("error"),
                    alert_sent,
                    json.dumps(daily_changes, ensure_ascii=False),
                ))
                conn.commit()
    except Exception as exc:
        print(f"Daily history save failed: {exc}")
    return daily_changes

def _history_load(limit: int = 100) -> list[dict[str, Any]]:
    if not _DATABASE_URL:
        return []
    try:
        with _db_connect() as conn:
            with conn.cursor() as cur:
                _db_init(cur)
                cur.execute("""
                    SELECT id, checked_at, started_at, finished_at, summary, error, alert_sent, changes, daily_date
                    FROM denni_pov_history
                    ORDER BY COALESCE(daily_date, checked_at::date) DESC, id DESC
                    LIMIT %s
                """, (max(1, min(limit, 500)),))
                rows = cur.fetchall()
                conn.commit()
        return [
            {
                "id": row[0],
                "checked_at": row[1].astimezone().strftime("%d.%m.%Y %H:%M:%S") if row[1] else "",
                "started_at": row[2] or "",
                "finished_at": row[3] or "",
                "summary": row[4] if isinstance(row[4], dict) else {},
                "error": row[5] or "",
                "alert_sent": bool(row[6]),
                "changes": row[7] if isinstance(row[7], dict) else {},
                "day": row[8].isoformat() if row[8] else "",
            }
            for row in rows
        ]
    except Exception as exc:
        print(f"History load failed: {exc}")
        return []


def _annotations_load_strict(legacy: dict[str, Any] | None = None) -> dict[str, dict[str, str]]:
    legacy_clean = _normalize_annotations(legacy)
    try:
        with _db_connect() as conn:
            with conn.cursor() as cur:
                _db_init(cur)

                # One-time / safety migration of any legacy in-state annotations.
                for vehicle_key, value in legacy_clean.items():
                    cur.execute("""
                        INSERT INTO denni_pov_annotations (vehicle_key, note, workflow_status, updated_at)
                        VALUES (%s, %s, %s, %s)
                        ON CONFLICT (vehicle_key) DO NOTHING
                    """, (
                        vehicle_key,
                        value.get("note", "")[:2000],
                        value.get("workflow_status", "")[:50],
                        value.get("updated_at") or _now(),
                    ))

                cur.execute("""
                    SELECT vehicle_key, note, workflow_status, updated_at
                    FROM denni_pov_annotations
                """)
                rows = cur.fetchall()
                conn.commit()

        return {
            str(row[0]).strip().upper(): {
                "note": row[1] or "",
                "workflow_status": row[2] or "",
                "updated_at": row[3] or "",
            }
            for row in rows
        }
    except PersistenceUnavailable:
        raise
    except Exception as exc:
        raise PersistenceUnavailable(_safe_db_error(exc)) from exc


def _annotations_load(legacy: dict[str, Any] | None = None) -> dict[str, dict[str, str]]:
    """Best-effort read for rendering; durable writes never use this fallback."""
    try:
        return _annotations_load_strict(legacy)
    except PersistenceUnavailable as exc:
        print(f"Annotations read fallback: {exc}")
        return _normalize_annotations(legacy)


def _annotation_save(vehicle_key: str, note: str, workflow_status: str, snapshot: dict[str, Any] | None = None) -> None:
    """Persist one vehicle annotation and an immutable audit record."""
    snapshot = snapshot if isinstance(snapshot, dict) else {}
    try:
        with _db_connect() as conn:
            with conn.cursor() as cur:
                _db_init(cur)
                cur.execute(
                    "SELECT note, workflow_status FROM denni_pov_annotations WHERE vehicle_key = %s",
                    (vehicle_key,),
                )
                old_row = cur.fetchone()
                old_note = str(old_row[0] or "") if old_row else ""
                old_status = str(old_row[1] or "") if old_row else ""

                if not note and not workflow_status:
                    cur.execute("DELETE FROM denni_pov_annotations WHERE vehicle_key = %s", (vehicle_key,))
                else:
                    cur.execute("""
                        INSERT INTO denni_pov_annotations (vehicle_key, note, workflow_status, updated_at)
                        VALUES (%s, %s, %s, %s)
                        ON CONFLICT (vehicle_key)
                        DO UPDATE SET
                            note = EXCLUDED.note,
                            workflow_status = EXCLUDED.workflow_status,
                            updated_at = EXCLUDED.updated_at
                    """, (vehicle_key, note, workflow_status, _now()))

                if old_note != note or old_status != workflow_status:
                    cur.execute("""
                        INSERT INTO denni_pov_annotation_history
                            (vehicle_key, vin, spz, vehicle, original_status,
                             old_workflow_status, new_workflow_status,
                             old_note, new_note, changed_at_text)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """, (
                        vehicle_key,
                        str(snapshot.get("vin") or ""),
                        str(snapshot.get("spz_tir") or snapshot.get("spz_uniqa") or ""),
                        str(snapshot.get("vozidlo") or ""),
                        str(snapshot.get("status_raw") or snapshot.get("status") or ""),
                        old_status,
                        workflow_status,
                        old_note,
                        note,
                        _now(),
                    ))
                conn.commit()
    except PersistenceUnavailable:
        raise
    except Exception as exc:
        raise PersistenceUnavailable(_safe_db_error(exc)) from exc


def _annotation_history_load(limit: int = 300) -> list[dict[str, Any]]:
    if not _DATABASE_URL:
        return []
    try:
        with _db_connect() as conn:
            with conn.cursor() as cur:
                _db_init(cur)
                cur.execute("""
                    SELECT vehicle_key, vin, spz, vehicle, original_status,
                           old_workflow_status, new_workflow_status,
                           old_note, new_note, changed_at_text
                    FROM denni_pov_annotation_history
                    ORDER BY id DESC
                    LIMIT %s
                """, (max(1, min(limit, 1000)),))
                rows = cur.fetchall()

                # Current annotations are also included even if they predate the audit table.
                cur.execute("""
                    SELECT vehicle_key, note, workflow_status, updated_at
                    FROM denni_pov_annotations
                """)
                current_rows = cur.fetchall()
                conn.commit()

        history = [{
            "vehicle_key": row[0] or "",
            "vin": row[1] or "",
            "spz": row[2] or "",
            "vehicle": row[3] or "",
            "original_status": row[4] or "",
            "old_workflow_status": row[5] or "",
            "new_workflow_status": row[6] or "",
            "old_note": row[7] or "",
            "new_note": row[8] or "",
            "changed_at": row[9] or "",
            "source": "history",
        } for row in rows]

        keys_in_history = {str(item.get("vehicle_key") or "") for item in history}
        for vehicle_key, note, workflow_status, updated_at in current_rows:
            if str(vehicle_key or "") not in keys_in_history:
                history.append({
                    "vehicle_key": vehicle_key or "",
                    "vin": vehicle_key if not str(vehicle_key or "").startswith("SPZ:") else "",
                    "spz": str(vehicle_key or "")[4:] if str(vehicle_key or "").startswith("SPZ:") else "",
                    "vehicle": "",
                    "original_status": "",
                    "old_workflow_status": "",
                    "new_workflow_status": workflow_status or "",
                    "old_note": "",
                    "new_note": note or "",
                    "changed_at": updated_at or "",
                    "source": "current",
                })
        history.sort(key=lambda item: str(item.get("changed_at") or ""), reverse=True)
        return history[:max(1, min(limit, 1000))]
    except Exception as exc:
        print(f"Annotation history load failed: {exc}")
        return []

def _file_load() -> dict[str, Any] | None:
    try:
        if _STATE_FILE.exists():
            data = json.loads(_STATE_FILE.read_text(encoding="utf-8"))
            if isinstance(data, dict):
                return data
    except Exception as exc:
        print(f"Cloud state file load failed: {exc}")
    return None


def _file_save(data: dict[str, Any]) -> None:
    try:
        _STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
        tmp = _STATE_FILE.with_suffix(".tmp")
        tmp.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        tmp.replace(_STATE_FILE)
    except Exception as exc:
        print(f"Cloud state file save failed: {exc}")


def _load_state() -> dict[str, Any]:
    data = _db_load() or _file_load() or _default_state()
    base = _default_state()
    base.update(data)
    if not isinstance(base.get("sources"), dict): base["sources"] = _default_state()["sources"]
    if not isinstance(base.get("summary"), dict): base["summary"] = _default_state()["summary"]
    if not isinstance(base.get("results"), list): base["results"] = []
    if not isinstance(base.get("changes"), dict): base["changes"] = _default_state()["changes"]
    if not isinstance(base.get("annotations"), dict): base["annotations"] = {}
    if not isinstance(base.get("progress"), dict): base["progress"] = _default_state()["progress"]
    return base


def _save_state(data: dict[str, Any]) -> None:
    if not _db_save(data):
        _file_save(data)


def _result_key(row: dict[str, Any]) -> str:
    vin = str(row.get("vin") or "").strip().upper()
    spz = str(row.get("spz_tir") or row.get("spz_uniqa") or "").strip().upper()
    return vin or f"SPZ:{spz}"


def _compare_runs(previous: dict[str, Any], current: dict[str, Any]) -> dict[str, Any]:
    prev_rows = {_result_key(r): r for r in (previous.get("results") or []) if _result_key(r)}
    curr_rows = {_result_key(r): r for r in (current.get("results") or []) if _result_key(r)}
    items: list[dict[str, Any]] = []

    def vehicle_payload(row: dict[str, Any]) -> dict[str, str]:
        return {
            "vin": str(row.get("vin") or ""),
            "spz": str(row.get("spz_tir") or row.get("spz_uniqa") or ""),
            "vozidlo": str(row.get("vozidlo") or ""),
        }

    for key, row in curr_rows.items():
        old = prev_rows.get(key)
        if old is None:
            items.append({"type": "NEW", "key": key, **vehicle_payload(row), "new_status": str(row.get("status_raw") or row.get("status") or "")})
            continue

        old_status = str(old.get("status_raw") or old.get("status") or "")
        new_status = str(row.get("status_raw") or row.get("status") or "")
        old_spz = str(old.get("spz_tir") or old.get("spz_uniqa") or "")
        new_spz = str(row.get("spz_tir") or row.get("spz_uniqa") or "")

        if old_status != new_status:
            items.append({
                "type": "STATUS",
                "key": key,
                **vehicle_payload(row),
                "old_status": old_status,
                "new_status": new_status,
            })
        if old_spz != new_spz:
            items.append({
                "type": "SPZ",
                "key": key,
                **vehicle_payload(row),
                "old_spz": old_spz,
                "new_spz": new_spz,
            })

    for key, row in prev_rows.items():
        if key not in curr_rows:
            items.append({"type": "REMOVED", "key": key, **vehicle_payload(row), "old_status": str(row.get("status_raw") or row.get("status") or "")})

    prev_summary = previous.get("summary") or {}
    curr_summary = current.get("summary") or {}
    summary_delta = {}
    for field in ("active", "ok_total", "missing", "absent_insured", "absent_uninsured", "deposit", "sold_uniqa", "extra_uniqa"):
        old_value = int(prev_summary.get(field) or 0)
        new_value = int(curr_summary.get(field) or 0)
        if old_value != new_value:
            summary_delta[field] = {"from": old_value, "to": new_value, "delta": new_value - old_value}

    return {
        "count": len(items),
        "items": items,
        "summary_delta": summary_delta,
        "compared_to": previous.get("finished_at") or previous.get("synced_at"),
    }


def _same_local_day(value: Any, day_iso: str) -> bool:
    return _date_from_display(value) == day_iso


def _public_state(data: dict[str, Any]) -> dict[str, Any]:
    public = dict(data)
    public.pop("_command", None)
    public.pop("annotations", None)
    annotations = _annotations_load(data.get("annotations"))
    rows = []
    for original in data.get("results") or []:
        row = dict(original)
        meta = annotations.get(_result_key(row), {})
        row["note"] = str(meta.get("note") or "")
        row["workflow_status"] = str(meta.get("workflow_status") or "")
        row["workflow_updated_at"] = str(meta.get("updated_at") or "")
        if row["note"]:
            row["original_detail"] = row.get("detail", "")
            row["detail"] = row["note"]
        if row["workflow_status"]:
            row["original_status"] = row.get("status", "")
            row["status"] = row["workflow_status"]
        rows.append(row)
    public["results"] = rows
    summary = dict(data.get("summary") or {})
    deposit = int(summary.get("deposit") or 0)
    summary["active"] = max(0, int(summary.get("active") or 0) - deposit)
    summary["ok_total"] = max(0, int(summary.get("ok_total") or 0) - deposit)

    resolved_statuses = {"VYŘEŠENO", "V POŘÁDKU"}
    issue_statuses = {
        "CHYBÍ V UNIQA",
        "NEPŘÍTOMNÉ, ALE POJIŠTĚNÉ",
        "PRODANÉ, ALE V UNIQA",
        "NAVÍC V UNIQA",
    }

    def _is_resolved_issue(row: dict[str, Any]) -> bool:
        raw = str(row.get("status_raw") or "").upper()
        workflow = str(row.get("workflow_status") or "").upper()
        if raw == "NEPŘÍTOMNÉ, ALE POJIŠTĚNÉ":
            return workflow == "VYŘEŠENO"
        return workflow in resolved_statuses

    resolved_issue_rows = [
        row for row in rows
        if str(row.get("status_raw") or "").upper() in issue_statuses
        and _is_resolved_issue(row)
    ]
    # Do počtu "Pojištění v pořádku" patří jen vyřešené problémy vozidel,
    # která jsou součástí aktivní kontroly. Záznamy "NAVÍC V UNIQA"
    # a prodaná vozidla nejsou aktivní flotila a nesmí zvyšovat ok_total.
    resolved_active_rows = [
        row for row in resolved_issue_rows
        if str(row.get("status_raw") or "").upper() in {
            "CHYBÍ V UNIQA",
            "NEPŘÍTOMNÉ, ALE POJIŠTĚNÉ",
        }
    ]
    summary["ok_total"] = min(
        int(summary.get("active") or 0),
        int(summary.get("ok_total") or 0) + len(resolved_active_rows),
    )
    summary["missing"] = sum(
        1 for row in rows
        if str(row.get("status_raw") or "").upper() == "CHYBÍ V UNIQA"
        and str(row.get("workflow_status") or "").upper() not in resolved_statuses
    )
    summary["absent_insured"] = sum(
        1 for row in rows
        if str(row.get("status_raw") or "").upper() == "NEPŘÍTOMNÉ, ALE POJIŠTĚNÉ"
        and str(row.get("workflow_status") or "").upper() != "VYŘEŠENO"
    )
    summary["sold_uniqa"] = sum(
        1 for row in rows
        if str(row.get("status_raw") or "").upper() == "PRODANÉ, ALE V UNIQA"
        and str(row.get("workflow_status") or "").upper() not in resolved_statuses
    )
    summary["extra_uniqa"] = sum(
        1 for row in rows
        if str(row.get("status_raw") or "").upper() == "NAVÍC V UNIQA"
        and str(row.get("workflow_status") or "").upper() not in resolved_statuses
    )
    summary["manual_ok"] = len(resolved_issue_rows)
    public["summary"] = summary
    public["csv_available"] = bool(rows)
    return public


def _authorized() -> bool:
    expected = os.getenv("SYNC_TOKEN", "")
    if not expected: return False
    supplied = request.headers.get("Authorization", "")
    return hmac.compare_digest(supplied, f"Bearer {expected}")


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/state")
def api_state():
    with _lock:
        data = _load_state()
    return jsonify(_public_state(data))


@app.get("/api/vehicle-technical/<vin>")
def api_vehicle_technical(vin: str):
    """Read the saved technical record for a vehicle in the dashboard."""
    vin = re.sub(r"\s+", "", vin).upper()
    if not _KOSTKA_VIN.fullmatch(vin):
        return jsonify({"ok": False, "message": "Neplatné VIN vozidla."}), 400

    with _lock:
        results = _load_state().get("results") or []
    if not any(re.sub(r"\s+", "", str(row.get("vin") or "")).upper() == vin
               for row in results if isinstance(row, dict)):
        return jsonify({"ok": False, "message": "Vozidlo není v aktuálním přehledu."}), 404

    try:
        saved = _kostka_saved(vin)
    except Exception:
        return jsonify({"ok": False, "message": "Uložené technické údaje teď nejsou dostupné."}), 503
    if saved:
        if saved[0]:
            return jsonify({"ok": True, "vin": vin, "data": saved[0], "fetched_at": saved[1]})
        return jsonify({"ok": True, "vin": vin, "message": "Datová kostka k tomuto VIN nemá technické údaje."})
    configured = bool(_kostka_keys() and _DATABASE_URL)
    message = ("Technické údaje se načítají na pozadí." if configured else
               "API Datové kostky není na serveru nastavené.")
    return jsonify({"ok": True, "vin": vin, "pending": configured, "message": message})


@app.get("/api/history")
def api_history():
    limit = request.args.get("limit", "100")
    try:
        limit_i = int(limit)
    except ValueError:
        limit_i = 100
    return jsonify({"ok": True, "history": _history_load(limit_i)})


@app.get("/api/manual-history")
def api_manual_history():
    limit = request.args.get("limit", "300")
    try:
        limit_i = int(limit)
    except ValueError:
        limit_i = 300
    return jsonify({"ok": True, "history": _annotation_history_load(limit_i)})


@app.post("/api/result-meta")
def api_result_meta():
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify({"ok": False, "persisted": False, "message": "Neplatná data."}), 400

    key = str(payload.get("key") or "").strip().upper()
    if not key:
        return jsonify({"ok": False, "persisted": False, "message": "Chybí identifikace vozidla."}), 400

    note = str(payload.get("note") or "").strip()[:2000]
    workflow_status = str(payload.get("workflow_status") or "").strip().upper()
    if workflow_status not in {"", "V POŘÁDKU", "VYŘEŠENO", "ŘEŠÍ SE", "KONTROLA"}:
        return jsonify({"ok": False, "persisted": False, "message": "Nepovolený status."}), 400

    # PostgreSQL is authoritative. Never claim success if this write fails.
    with _lock:
        current_data = _load_state()
        snapshot = next(
            (dict(row) for row in (current_data.get("results") or []) if _result_key(row) == key),
            {},
        )
    try:
        _annotation_save(key, note, workflow_status, snapshot)
    except PersistenceUnavailable as exc:
        print(f"Annotation durable save failed for {key}: {exc}")
        return jsonify({
            "ok": False,
            "persisted": False,
            "storage": "postgres",
            "message": "Status ani poznámka NEBYLY uloženy. Trvalé úložiště PostgreSQL není dostupné.",
            "detail": str(exc),
        }), 503

    # Keep an in-state copy only as a migration/cache aid. It is not the source of truth.
    with _lock:
        data = _load_state()
        annotations = data.setdefault("annotations", {})
        if note or workflow_status:
            annotations[key] = {
                "note": note,
                "workflow_status": workflow_status,
                "updated_at": _now(),
            }
        else:
            annotations.pop(key, None)
        _save_state(data)

    return jsonify({
        "ok": True,
        "persisted": True,
        "storage": "postgres",
        "message": "Status a poznámka byly trvale uloženy a jsou sdílené mezi počítači.",
    })

@app.post("/api/run")
def api_run():
    with _lock:
        data = _load_state()
        if data.get("running"):
            return jsonify({"ok": False, "message": "Kontrola už probíhá."}), 409
        command_id = uuid.uuid4().hex
        data["running"] = True
        data["started_at"] = _now()
        data["finished_at"] = None
        data["error"] = None
        data["progress"] = {"percent": 3, "phase": "Čekám na kancelářský agent", "eta_seconds": 90}
        data["_command"] = {"id": command_id, "action": "run_check", "requested_at": _now()}
        data["sources"] = {
            "tirbazar": {"state": "loading", "status": "Čekám na kancelářský agent…", "detail": "TIRBazar SQL je dostupný pouze z interní sítě"},
            "uniqa": {"state": "idle", "status": "Čekám…", "detail": "AIV / Denní POV / Aktivní"},
            "allianz": {"state": "idle", "status": "Čekám…", "detail": "Flotilové PDF"},
        }
        _save_state(data)
    return jsonify({"ok": True, "message": "Požadavek na kontrolu odeslán kancelářskému agentovi."})


@app.get("/api/agent/command")
def agent_command():
    if not _authorized(): return jsonify({"ok": False, "message": "Unauthorized"}), 401
    with _lock:
        data = _load_state()
        command = data.get("_command")
    return jsonify({"ok": True, "command": command})


@app.post("/api/sync")
def api_sync():
    if not _authorized(): return jsonify({"ok": False, "message": "Unauthorized"}), 401
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict): return jsonify({"ok": False, "message": "Neplatná data synchronizace."}), 400
    required = ("sources", "summary", "results")
    if any(key not in payload for key in required): return jsonify({"ok": False, "message": "Synchronizace nemá všechny povinné části."}), 400
    if not isinstance(payload.get("sources"), dict) or not isinstance(payload.get("summary"), dict) or not isinstance(payload.get("results"), list):
        return jsonify({"ok": False, "message": "Neplatný formát synchronizace."}), 400

    final_run = not bool(payload.get("running"))
    alert_sent = False

    with _lock:
        previous = _load_state()
        annotations = _annotations_load(previous.get("annotations"))
        data = _default_state()
        data["annotations"] = annotations
        for key in ("running", "started_at", "finished_at", "error", "sources", "summary", "results", "progress"):
            if key in payload: data[key] = payload[key]
        data["running"] = bool(payload.get("running"))
        if data["running"]:
            data["finished_at"] = None
        else:
            data["finished_at"] = payload.get("finished_at") or _now()
            if not data.get("error"):
                data["progress"] = {"percent": 100, "phase": "Hotovo", "eta_seconds": 0}
        data["synced_at"] = _now()
        data["csv_available"] = bool(data.get("results"))
        data["_command"] = None
        if final_run and previous.get("finished_at") and previous.get("results"):
            run_changes = _compare_runs(previous, data)
            if _same_local_day(previous.get("finished_at"), _today_iso()):
                prior_daily = previous.get("changes") if isinstance(previous.get("changes"), dict) else {}
                merged_items = _merge_change_items(prior_daily.get("items") or [], run_changes.get("items") or [])
                data["changes"] = {
                    "count": len(merged_items),
                    "items": merged_items,
                    "summary_delta": run_changes.get("summary_delta") or prior_daily.get("summary_delta") or {},
                    "compared_to": prior_daily.get("compared_to") or run_changes.get("compared_to"),
                    "day": _today_iso(),
                }
            else:
                run_changes["day"] = _today_iso()
                data["changes"] = run_changes
        elif final_run:
            data["changes"] = {"count": 0, "items": [], "summary_delta": {}, "compared_to": previous.get("finished_at") or previous.get("synced_at"), "day": _today_iso()}
        else:
            data["changes"] = previous.get("changes") or _default_state()["changes"]

        if final_run:
            data["changes"] = _history_save_daily(data, alert_sent)
        _save_state(data)

    if final_run and not data.get("error"):
        _kostka_start(data.get("results") or [])

    return jsonify({"ok": True, "message": "Data byla synchronizována na Render.", "alert_sent": alert_sent})



def _excel_text(value: Any) -> str:
    text = str(value or "")
    return re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", text)


def _build_xlsx_report(data: dict[str, Any]) -> bytes:
    rows = list(data.get("results") or [])
    summary = dict(data.get("summary") or {})

    wb = Workbook()
    ws = wb.active
    ws.title = "Přehled pojištění"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A11"

    navy = "0D2740"
    blue = "1684D6"
    green = "169B63"
    red = "D9143B"
    red_soft = "FFF0F3"
    green_soft = "E9F7F0"
    blue_soft = "EAF5FD"
    gray = "6C8195"
    light = "F5F8FA"
    border_color = "DCE6ED"
    white = "FFFFFF"
    dark = "17354E"

    thin = Side(style="thin", color=border_color)

    ws.merge_cells("A1:H2")
    c = ws["A1"]
    c.value = "DENNÍ POV  |  VANS CENTRE"
    c.font = Font(name="Arial", size=22, bold=True, color=white)
    c.fill = PatternFill(fill_type="solid", fgColor=navy)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 26

    ws.merge_cells("A3:H3")
    c = ws["A3"]
    c.value = f"Přehled kontroly pojištění vozidel  •  {data.get('finished_at') or data.get('synced_at') or _now()}"
    c.font = Font(name="Arial", size=10, color=gray)
    c.fill = PatternFill(fill_type="solid", fgColor="F8FBFD")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 22

    cards = [
        ("A5:B5", "A6:B7", "AKTIVNÍ VOZIDLA", int(summary.get("active") or 0), blue, blue_soft),
        ("C5:D5", "C6:D7", "POJIŠTĚNÍ V POŘÁDKU", int(summary.get("ok_total") or 0), green, green_soft),
        ("E5:F5", "E6:F7", "CHYBÍ POJIŠTĚNÍ", int(summary.get("missing") or 0), red, red_soft),
        ("G5:H5", "G6:H7", "NEPŘÍTOMNÉ · POJIŠTĚNO", int(summary.get("absent_insured") or 0), red, red_soft),
    ]
    for label_rng, value_rng, label, value, accent, soft in cards:
        ws.merge_cells(label_rng)
        ws.merge_cells(value_rng)
        lc = ws[label_rng.split(":")[0]]
        vc = ws[value_rng.split(":")[0]]
        lc.value = label
        lc.font = Font(name="Arial", size=9, bold=True, color=gray)
        lc.fill = PatternFill(fill_type="solid", fgColor=soft)
        lc.alignment = Alignment(horizontal="center", vertical="center")
        vc.value = value
        vc.font = Font(name="Arial", size=24, bold=True, color=accent)
        vc.fill = PatternFill(fill_type="solid", fgColor=soft)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws[label_rng]:
            for cell in row:
                cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
        for row in ws[value_rng]:
            for cell in row:
                cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)

    ws.merge_cells("A9:H9")
    ws["A9"] = "DETAILNÍ PŘEHLED VOZIDEL"
    ws["A9"].font = Font(name="Arial", size=11, bold=True, color=dark)
    ws["A9"].alignment = Alignment(vertical="center")
    ws.row_dimensions[9].height = 24

    headers = ["Stav", "VIN", "SPZ", "Datum výkupu", "Datum prodeje", "Výsledek", "Poznámka", "Řešení"]
    for col, value in enumerate(headers, start=1):
        cell = ws.cell(row=10, column=col, value=value)
        cell.font = Font(name="Arial", size=9, bold=True, color=white)
        cell.fill = PatternFill(fill_type="solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = Border(bottom=thin)
    ws.row_dimensions[10].height = 24

    status_fills = {
        "OK": ("E9F7F0", green),
        "NEPŘÍTOMNÉ, ALE NEPOJIŠTĚNÉ": ("E9F7F0", green),
        "CHYBÍ V UNIQA": ("FFF0F3", red),
        "NEPŘÍTOMNÉ, ALE POJIŠTĚNÉ": ("FFF0F3", red),
        "PRODANÉ, ALE V UNIQA": ("F4F0FF", "7256B8"),
        "NEPOJIŠTĚNO, ALE DEPOZIT": ("FFF7E7", "A87512"),
        "NAVÍC V UNIQA": ("EAF7FA", "16849B"),
        "SPZ NESOUHLASÍ": ("FFF7E7", "A87512"),
        "NELZE OVĚŘIT": ("F1F4F6", "647789"),
    }

    start_row = 11
    for idx, row in enumerate(rows, start=start_row):
        values = [
            row.get("status") or row.get("status_raw") or "",
            row.get("vin") or "",
            row.get("spz_tir") or row.get("spz_uniqa") or "",
            row.get("vykup") or "",
            row.get("prodej") or "",
            row.get("detail") or "",
            row.get("note") or "",
            row.get("workflow_status") or "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=idx, column=col, value=_excel_text(value))
            cell.font = Font(name="Arial", size=9, color=dark)
            cell.alignment = Alignment(vertical="top", wrap_text=(col in {6, 7}))
            cell.border = Border(bottom=Side(style="hair", color="E5EDF2"))
            if idx % 2 == 0:
                cell.fill = PatternFill(fill_type="solid", fgColor="FBFDFE")

        raw = str(row.get("status_raw") or "").upper()
        fill_color, font_color = status_fills.get(raw, ("F1F4F6", "647789"))
        ws.cell(row=idx, column=1).fill = PatternFill(fill_type="solid", fgColor=fill_color)
        ws.cell(row=idx, column=1).font = Font(name="Arial", size=9, bold=True, color=font_color)
        ws.row_dimensions[idx].height = 30 if any(values[5:7]) else 22

    end_row = max(10, start_row + len(rows) - 1)
    ws.auto_filter.ref = f"A10:H{end_row}"

    widths = {"A": 28, "B": 22, "C": 14, "D": 15, "E": 15, "F": 56, "G": 38, "H": 16}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    footer_row = end_row + 2
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=8)
    fc = ws.cell(row=footer_row, column=1)
    fc.value = f"Vygenerováno: {_now()}  •  Denní POV / Vans Centre"
    fc.font = Font(name="Arial", size=8, color=gray)
    fc.alignment = Alignment(horizontal="right")

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.print_title_rows = "1:10"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


@app.get("/download/xlsx")
def download_xlsx():
    with _lock:
        data = _public_state(_load_state())
        rows = list(data.get("results") or [])
    if not rows:
        return jsonify({"ok": False, "message": "Přehled zatím není k dispozici."}), 404

    body = _build_xlsx_report(data)
    filename = f"denni_pov_prehled_{datetime.now().strftime('%Y-%m-%d_%H-%M')}.xlsx"
    return Response(
        body,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/download/csv")
def download_csv():
    with _lock:
        data = _public_state(_load_state())
        rows = list(data.get("results") or [])
    if not rows: return jsonify({"ok": False, "message": "CSV zatím není k dispozici."}), 404
    stream = io.StringIO()
    writer = csv.writer(stream, delimiter=";")
    writer.writerow(["Stav", "Pojišťovna", "VIN", "SPZ", "Datum výkupu", "Datum prodeje", "Výsledek", "Poznámka"])
    for row in rows:
        writer.writerow([row.get("status", ""), row.get("pojistovna", ""), row.get("vin", ""), row.get("spz_tir") or row.get("spz_uniqa") or "", row.get("vykup", ""), row.get("prodej", ""), row.get("detail", ""), row.get("note", "")])
    body = "\ufeff" + stream.getvalue()
    return Response(body, mimetype="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename=denni_pov_online.csv"})


@app.get("/health")
def health():
    with _lock:
        data = _load_state()

    db_ready, db_error = _db_probe()
    return jsonify({
        "ok": True,
        "service": "DENNI POV - KONTROLA",
        "mode": "online",
        "storage": "postgres" if db_ready else "file-fallback",
        "persistence": {
            "ready": db_ready,
            "configured": bool(_DATABASE_URL),
            "annotations": "postgres" if db_ready else "unavailable",
            "error": db_error,
        },
        "synced_at": data.get("synced_at"),
        "waiting_for_agent": bool(data.get("_command")),
    })
