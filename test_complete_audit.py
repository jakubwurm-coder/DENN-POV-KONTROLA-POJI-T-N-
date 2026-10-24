import copy
import io
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch
from openpyxl import load_workbook
from models import TirVehicle, UniqaVehicle
from compare import compare_vehicles
from audit_validation import validate_complete_audit
from xlsx_report import write_report
import tirbazar
import local_app


def vehicle(oid=1, **changes):
    return TirVehicle(**dict(dict(oid=oid, vin='VIN'+str(oid), spz='1AA1111',
        zeme_puvodu='CZ', stav='VYKOUPENÉ', datum_vykupu='2026-01-01'), **changes))


def compare(rows, uq=(), al=(), available=True):
    return compare_vehicles(rows, [UniqaVehicle(v, '1AA1111') for v in uq], True,
        allianz=[SimpleNamespace(vin=v, spz='') for v in al], allianz_available=available,
        known_tir_vins={r.vin for r in rows if r.vin}, complete=True)


class CompleteComparisonTests(TestCase):
    def test_all_states_and_dual_have_one_result_each(self):
        states=['PRODANÉ','VYKOUPENÉ','NEPŘÍTOMNÉ','V KOMISI','REZERVOVANÉ',
            'PRONAJATÉ','VOLNÉ','PARKOVANÉ','PARKOVÁNÍ UKONČENO','VRÁCENÉ Z KOMISE','NOVÝ STAV','']
        rows=[vehicle(i, stav=s) for i,s in enumerate(states, 1)]
        vins=[v.vin for v in rows]
        results=compare(rows,vins,vins)
        self.assertEqual(len(results),len(rows))
        self.assertEqual({r.status for r in results},{'DVOJÍ POJIŠTĚNÍ'})
        self.assertEqual(len(compare(rows)),len(rows))

    def test_sold_no_age_cutoff_and_neutral_audit(self):
        sold=vehicle(stav='PRODANÉ', datum_prodeje='2020-01-01')
        self.assertTrue(tirbazar._requires_pov_check(sold))
        self.assertEqual(compare([sold],['VIN1'])[0].status,'PRODANÉ, ALE POJIŠTĚNÉ')
        self.assertEqual(compare([sold])[0].status,'PRODANÉ, NEPOJIŠTĚNÉ')
        self.assertEqual(compare([sold],available=False)[0].status,'NELZE OVĚŘIT')

    def test_physical_duplicate_and_blank_vin(self):
        rows=[vehicle(),vehicle(2,vin='VIN1',stav='PRODANÉ'),vehicle(3,vin='')]
        results=compare(rows,['VIN1'])
        self.assertEqual({r.oid for r in results},{1,2,3})
        self.assertEqual(next(r.status for r in results if r.oid==3),'NELZE OVĚŘIT')

    def test_insurance_only_dual_once(self):
        results=compare([],['OUTSIDE'],['OUTSIDE'])
        self.assertEqual(len(results),1)
        self.assertEqual(results[0].status,'DVOJÍ POJIŠTĚNÍ')

    def test_worker_export_independent_of_visible_results(self):
        rows=[vehicle(), vehicle(2,vin='VIN1',stav='PRODANÉ',datum_prodeje='2020-01-01'),
            vehicle(3,vin=''),vehicle(4,stav='PRODANÉ',datum_prodeje='2020-01-01'),
            vehicle(5,stav='PRONAJATÉ')]
        selected=[rows[0],rows[3],rows[4]]
        original=copy.deepcopy(local_app._state)
        self.addCleanup(lambda: (local_app._state.clear(),local_app._state.update(original)))
        uq=SimpleNamespace(vehicles=[UniqaVehicle('VIN1','1AA1111')],
            raw_vehicles=[UniqaVehicle('VIN1','1AA1111',cps='A'),UniqaVehicle('VIN1','1AA1111',cps='B')], available=True,error='',duplicates=[])
        al=SimpleNamespace(vehicles=[],available=True,error='')
        with patch.object(local_app,'load_config'), patch.object(local_app,'load_tirbazar_vehicles',return_value=(selected,[[rows[0],rows[1]]])), \
             patch.object(tirbazar,'LAST_RAW_ROWS',rows),patch.object(tirbazar,'LAST_LOAD_STATS',{'sql_total':5}), \
             patch.object(local_app,'load_uniqa_vehicles',return_value=uq),patch.object(local_app,'load_allianz_vehicles',return_value=al), \
             patch.object(local_app,'prepare_output'),patch.object(local_app,'write_tirbazar_snapshot'), \
             patch.object(local_app,'write_duplicates'),patch.object(local_app,'write_comparison',return_value='/tmp/unused.csv'):
            local_app._run_check_worker()
        self.assertIsNone(local_app._state['error'])
        audit=local_app._state['audit']
        validate_complete_audit(audit)
        self.assertEqual(len(audit['sql']),5)
        self.assertEqual(len(audit['results']),5)
        self.assertEqual(len(audit['uniqa']),2)
        self.assertEqual([r.oid for r in local_app._state['results']],[1])
        self.assertEqual(local_app._state['results'][0].status,'DVOJÍ POJIŠTĚNÍ')
        output=io.BytesIO()
        write_report({'results':[], 'audit':audit},output)
        book=load_workbook(io.BytesIO(output.getvalue()))
        self.assertEqual(book['SQL audit'].max_row,6)
        self.assertEqual({book['Přehled pojištění'].cell(r,9).value for r in range(11,16)}, {'1','2','3','4','5'})
        self.assertEqual(book['Uniqa'].max_row,3)
        for bad in ({**audit,'results':audit['results'][:-1]}, {**audit,'schema_version':1},
                    {**audit,'sql':audit['sql'][:-1]}):
            with self.assertRaises(ValueError): validate_complete_audit(bad)

    def test_read_only_guard(self):
        tirbazar.validate_read_only(tirbazar.build_sql())
        for command in tirbazar.FORBIDDEN_SQL:
            with self.assertRaises(RuntimeError): tirbazar.validate_read_only(command+'dbo.Vozidlo')

class AuditApiTests(TestCase):
    def test_complete_results_are_saved_with_notes_and_do_not_replace_web_rows(self):
        import cloud_app as cloud
        results=[{'oid':i,'vin':'SAME' if i<3 else '', 'status_raw':'OK'} for i in (1,2,3)]
        audit={'schema_version':2, 'sql':[{'oid':i} for i in (1,2,3)],
               'results':results,'counts':{'sql_total':3,'physical_audit_rows':3}}
        state={'_audit_export':{'id':'audit'}, '_command':{'id':'audit'},
               'results':[results[0]], 'annotations':{}}
        snapshot={'finished_at':'24.10.2026 23:00:00','results':[results[0]]}
        saved={}
        with patch.object(cloud,'_authorized',return_value=True), \
             patch.object(cloud,'_load_state',return_value=state), \
             patch.object(cloud,'_save_state',side_effect=lambda s:saved.update(copy.deepcopy(s))), \
             patch.object(cloud,'_annotations_load_strict',return_value={'SAME':{'note':'Poznámka'}}):
            response=cloud.app.test_client().post('/api/agent/audit-result',json={
                'id':'audit','audit':audit,'snapshot':snapshot})
            self.assertEqual(response.status_code,200,response.json)
            self.assertEqual(len(saved['results']),1)
            export=saved['_audit_export_snapshot']
            self.assertEqual(len(export['results']),3)
            self.assertEqual(export['results'][0]['note'],'Poznámka')
            self.assertEqual(export['audit']['results'],export['results'])
            audit['results']=results[:1]
            rejected=cloud.app.test_client().post('/api/agent/audit-result',json={
                'id':'audit','audit':audit,'snapshot':snapshot})
            self.assertEqual(rejected.status_code,409)
