import io
import json
import os
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch
from openpyxl import load_workbook

import cloud_app as cloud
import export_jobs
from xlsx_report import write_report


FIXTURE = {
    'results': [{'vin': 'VIN-A', 'status': 'OK', 'note': '=1+1'}],
    'audit': {
        'sql': [{'oid': '1', 'vin': 'SAME', 'duvod': 'aktivní'},
                {'oid': '2', 'vin': 'SAME', 'duvod': 'starší karta'},
                {'oid': '3', 'vin': '', 'duvod': 'bez VIN'}],
        'uniqa': [{'vin': 'SAME'}], 'allianz': [{'vin': 'SAME'}],
        'breakdown': {'decisions': {'aktivní': 1}},
        'duplicates': [[{'oid': '1'}, {'oid': '2'}]],
        'counts': {'sql_total': 3, 'physical_audit_rows': 3},
    },
}


class ReportTests(unittest.TestCase):
    def test_all_physical_rows_and_sheets_and_measured_progress(self):
        output = io.BytesIO()
        progress = []
        write_report(FIXTURE, output, lambda *values: progress.append(values))
        book = load_workbook(io.BytesIO(output.getvalue()))
        self.assertEqual(book.sheetnames, ['Přehled pojištění', 'SQL audit', 'Rozpad TIRBazar', 'Uniqa', 'Allianz', 'Duplicity a výjimky'])
        self.assertEqual(book['SQL audit'].max_row, 4)
        self.assertEqual([book['SQL audit'].cell(r, 4).value for r in range(2, 5)], ['1', '2', '3'])
        self.assertEqual(book['Přehled pojištění']['G11'].value, '=1+1')
        self.assertEqual(book['Přehled pojištění']['G11'].data_type, 's')
        self.assertEqual(progress[-1], (9, 9, 'finalizing'))
        self.assertEqual([p[0] for p in progress if p[2] == 'rows'], list(range(10)))


class ExportApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = patch.object(export_jobs, 'ROOT', Path(self.temp.name))
        root.start(); self.addCleanup(root.stop)
        config = patch.dict(cloud.app.config, TESTING=True, SECRET_KEY='test-only')
        config.start(); self.addCleanup(config.stop)
        self.client = cloud.app.test_client()
        self.other = cloud.app.test_client()
        for client in (self.client, self.other):
            with client.session_transaction(base_url='https://localhost') as session:
                session['denni_pov_authenticated'] = True

    def get(self, path, client=None):
        return (client or self.client).get(path, base_url='https://localhost')

    def start(self, data=FIXTURE):
        with patch.object(cloud, '_load_state', return_value=data), patch.object(cloud, '_annotations_load', return_value={}):
            return self.client.post('/api/xlsx-export', base_url='https://localhost')

    def wait(self, job_id):
        deadline = time.monotonic() + 15
        seen = []
        while time.monotonic() < deadline:
            response = self.get('/api/xlsx-export/' + job_id)
            self.assertEqual(response.status_code, 200)
            seen.append(response.json)
            if response.json['status'] != 'running':
                return response.json
            time.sleep(.05)
        self.fail('Worker did not finish')

    def test_real_background_process_download_length_range_and_ownership(self):
        response = self.start()
        self.assertEqual(response.status_code, 202)
        self.assertNotIn('owner', response.json)
        job_id = response.json['id']
        status = self.wait(job_id)
        self.assertEqual(status['status'], 'ready', status)
        self.assertEqual(status['completed'], status['total'])
        self.assertFalse((Path(self.temp.name) / job_id / 'snapshot.json').exists())
        self.assertEqual(self.get('/api/xlsx-export/' + job_id, self.other).status_code, 404)
        self.assertEqual(self.get('/download/xlsx?job=' + job_id, self.other).status_code, 404)
        response = self.get('/download/xlsx?job=' + job_id)
        body = response.get_data()
        self.assertEqual(len(body), status['size'])
        self.assertEqual(int(response.headers['Content-Length']), len(body))
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        response.close()
        self.assertEqual(load_workbook(io.BytesIO(body))['SQL audit'].max_row, 4)
        partial = self.client.get('/download/xlsx?job=' + job_id, base_url='https://localhost', headers={'Range': 'bytes=0-31'})
        self.assertEqual(partial.status_code, 206)
        self.assertEqual(partial.get_data(), body[:32])
        partial.close()
        self.assertEqual(self.get('/download/xlsx').status_code, 200)

    def test_login_and_origin_are_enforced(self):
        anonymous = cloud.app.test_client()
        self.assertEqual(anonymous.post('/api/xlsx-export', base_url='https://localhost').status_code, 401)
        self.assertEqual(self.get('/api/xlsx-export/' + 'a'*32, anonymous).status_code, 401)
        self.assertEqual(self.get('/download/xlsx?job=' + 'a'*32, anonymous).status_code, 401)
        self.assertEqual(self.client.post('/api/xlsx-export', base_url='https://localhost', headers={'Origin':'https://evil.example'}).status_code, 403)
        self.assertEqual(self.get('/api/xlsx-export/not-a-job').status_code, 404)

    def test_empty_state(self):
        self.assertEqual(self.start({}).status_code, 404)

    def test_duplicate_start_reuses_job_and_other_session_cannot_queue(self):
        with patch.object(export_jobs.subprocess, 'Popen') as popen:
            popen.return_value.pid = os.getpid()
            first = self.start().json
            second = self.start().json
            self.assertEqual(first['id'], second['id'])
            self.assertEqual(popen.call_count, 1)
            other = self.other.post('/api/xlsx-export', base_url='https://localhost')
            self.assertEqual(other.status_code, 409)
            pending = self.get('/download/xlsx?job=' + first['id'])
            self.assertEqual(pending.status_code, 409)

    def test_interrupted_expired_failed_jobs(self):
        job_id = 'a'*32
        folder = Path(self.temp.name) / job_id
        folder.mkdir()
        with self.client.session_transaction(base_url='https://localhost') as session:
            session['xlsx_owner'] = 'test-owner'
        status = dict(id=job_id, owner='test-owner', created=time.time(), status='running')
        export_jobs.atomic_json(folder/'status.json', status)
        self.assertEqual(self.get('/api/xlsx-export/' + job_id).json['status'], 'error')
        status.update(created=time.time()-export_jobs.TTL-1)
        export_jobs.atomic_json(folder/'status.json', status)
        self.assertEqual(self.get('/api/xlsx-export/' + job_id).json['status'], 'expired')
        status.update(created=time.time(), status='running')
        export_jobs.atomic_json(folder/'status.json', status)
        (folder/'snapshot.json').write_text('invalid json')
        with patch('logging.exception'):
            export_jobs.run_worker(folder)
        self.assertEqual(json.loads((folder/'status.json').read_text())['status'], 'error')
        self.assertFalse((folder/'snapshot.json').exists())


if __name__ == '__main__':
    unittest.main()
