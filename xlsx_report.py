"""Sequential XLSX writer. No SQL access and no retained worksheet cell graph."""
from __future__ import annotations
import re
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

AUDIT_SHEETS = [('SQL audit',
  'sql_rows',
  [('kontrolovano', 'Kontrolováno'),
   ('zarazeni', 'Zařazení'),
   ('duvod', 'Důvod rozhodnutí'),
   ('oid', 'OID'),
   ('vin', 'VIN'),
   ('spz', 'SPZ'),
   ('stav', 'Stav'),
   ('zeme_puvodu', 'Země'),
   ('datum_vykupu', 'Datum výkupu'),
   ('datum_prodeje', 'Datum prodeje'),
   ('ocekavani', 'Očekávání'),
   ('uniqa', 'Uniqa'),
   ('allianz', 'Allianz'),
   ('vysledek', 'Výsledek'),
   ('filtr', 'Důvod filtru'), ('status_raw', 'Výsledek porovnání'),
   ('detail', 'Detail porovnání'), ('zdroj_udaju', 'Zdroj údajů')]),
 ('Rozpad TIRBazar',
  'breakdown_rows',
  [('skupina', 'Skupina'), ('polozka', 'Kategorie / důvod'), ('pocet', 'Počet')]),
 ('Uniqa',
  'uniqa_rows',
  [('vin', 'VIN'),
   ('spz', 'SPZ'),
   ('cps', 'ČPS'),
   ('poj_od', 'Pojištění od'),
   ('poj_do', 'Pojištění do')]),
 ('Allianz',
  'allianz_rows',
  [('identifikator', 'Identifikátor'), ('vin', 'VIN'), ('spz', 'SPZ'),
   ('pojistka', 'Pojistka'),
   ('poj_od', 'Pojištění od'),
   ('poj_do', 'Pojištění do')]),
 ('Duplicity a výjimky',
  'duplicate_rows',
  [('skupina', 'Skupina'),
   ('oid', 'OID'),
   ('vin', 'VIN'),
   ('spz', 'SPZ'),
   ('stav', 'Stav'),
   ('datum_vykupu', 'Datum výkupu'),
   ('datum_prodeje', 'Datum prodeje'),
   ('ocekavani', 'Očekávání')])]


def write_report(data, destination, progress=lambda *args: None):
    audit = data.get("audit") or {}
    breakdown = audit.get("breakdown") or {}
    breakdown_rows = [
        {"skupina": label, "polozka": name, "pocet": count}
        for key, label in (("decisions", "Kategorie rozhodnutí"),
                           ("states_active", "Stav – aktivní POV"),
                           ("reasons", "Důvod zařazení / vyřazení"))
        for name, count in sorted((breakdown.get(key) or {}).items())
    ]
    duplicate_rows = [
        {"skupina": group_no, **item}
        for group_no, group in enumerate(audit.get("duplicates") or [], 1)
        if isinstance(group, list) for item in group if isinstance(item, dict)
    ]
    sources = {"sql_rows": audit.get("sql") or [], "uniqa_rows": audit.get("uniqa") or [],
               "allianz_rows": audit.get("allianz") or [], "breakdown_rows": breakdown_rows,
               "duplicate_rows": duplicate_rows}
    rows = audit.get("results") if "results" in audit else data.get("results") or []
    total = len(rows) + sum(len(value) for value in sources.values())
    done = 0
    wb = Workbook(write_only=True)
    font = Font(name="Arial", size=10, color="17354E")
    heading = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0D2740")
    alignment = Alignment(vertical="top", wrap_text=True)

    def append(sheet, values, header=False):
        cells = []
        for value in values:
            # Always literal strings: external SQL/provider values must never become formulas.
            cell = WriteOnlyCell(sheet, value=re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", str(value if value is not None else "")))
            cell.data_type = "s"
            cell.font = heading if header else font
            cell.alignment = alignment
            if header: cell.fill = fill
            cells.append(cell)
        sheet.append(cells)

    sheet = wb.create_sheet("Přehled pojištění")
    sheet.freeze_panes = "A11"
    sheet.sheet_view.showGridLines = False
    for col, width in enumerate((28, 22, 14, 15, 15, 56, 38, 16, 14, 24), 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    summary = data.get("summary") or {}
    append(sheet, ["DENNÍ POV | VANS CENTRE"], True)
    append(sheet, ["Přehled kontroly pojištění vozidel", data.get("finished_at") or data.get("synced_at") or ""])
    append(sheet, [])
    for label, key in (("AKTIVNÍ VOZIDLA", "active"), ("POJIŠTĚNÍ V POŘÁDKU", "ok_total"),
                       ("CHYBÍ POJIŠTĚNÍ", "missing"), ("NEPŘÍTOMNÉ · POJIŠTĚNO", "absent_insured")):
        append(sheet, [label, summary.get(key, 0)])
    append(sheet, [])
    append(sheet, ["DETAILNÍ PŘEHLED VOZIDEL"], True)
    append(sheet, ["Stav", "VIN", "SPZ", "Datum výkupu", "Datum prodeje", "Výsledek", "Poznámka", "Řešení", "OID", "Zdroj údajů"], True)
    progress(done, total, "rows")
    for row in rows:
        append(sheet, [row.get("status") or row.get("status_raw"), row.get("vin"),
                       row.get("spz_tir") or row.get("spz_uniqa"), row.get("vykup"), row.get("prodej"),
                       row.get("detail"), row.get("note"), row.get("workflow_status"), row.get("oid"),
                       "SQL Vans Centre" if row.get("oid") is not None else row.get("technical_source") or "UNIQA / Allianz"])
        done += 1
        progress(done, total, "rows")
    sheet.auto_filter.ref = f"A10:J{max(10, 10 + len(rows))}"
    append(sheet, [])
    append(sheet, ["AUDIT KONTROLY"], True)
    for label, key in (("SQL záznamů celkem", "sql_total"), ("Řádků v auditu OID", "physical_audit_rows"),
                       ("Bez VIN – audit", "without_vin_audit"), ("Duplicitní SQL karty", "duplicate_audit_rows"),
                       ("Unikátních VIN", "unique_vins"), ("Má být pojištěno", "expected_insured"),
                       ("Má být nepojištěno", "expected_uninsured"), ("Odfiltrováno", "filtered"),
                       ("Uniqa záznamů", "uniqa"), ("Allianz záznamů", "allianz"),
                       ("Duplicitních VIN skupin", "duplicate_vin_groups")):
        append(sheet, [label, (audit.get("counts") or {}).get(key, "")])

    for title, source, columns in AUDIT_SHEETS:
        sheet = wb.create_sheet(title)
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for col, (key, label) in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(col)].width = min(42, max(12, len(label) + 3, 20 if key in {"vin", "vysledek", "filtr"} else 12))
        append(sheet, [label for _, label in columns], True)
        for item in sources[source]:
            append(sheet, [item.get(key, "") for key, _ in columns])
            done += 1
            progress(done, total, "rows")
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(sources[source]) + 1}"
    # ZIP serialization has no reliable work denominator. Do not fabricate a percentage.
    progress(done, total, "finalizing")
    wb.save(destination)
